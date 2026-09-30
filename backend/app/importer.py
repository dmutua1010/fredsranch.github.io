"""Admin Excel ingestion. Sheets: Positions, KPIs, Records. Everything is validated first;
if any row fails NOTHING is written and row-level errors are returned (use dry_run to preview)."""
import hashlib, io
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from .models import Position, KPI, Record, ImportBatch

SHEETS = {
    "Positions": {"required": ["code", "title"]},
    "KPIs": {"required": ["code", "position_code", "kra", "name", "target"]},
    "Records": {"required": ["kpi_code", "period"]},
}

def _rows(ws, spec, errors):
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    missing = [c for c in spec["required"] if c not in header]
    if missing:
        errors.append(f"{ws.title}: missing column(s) {missing}")
        return []
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        d = {h: (row[j] if j < len(row) else None) for j, h in enumerate(header) if h}
        for c in spec["required"]:
            if d.get(c) in (None, ""):
                errors.append(f"{ws.title} row {i}: '{c}' is required")
        out.append((i, d))
    return out

def _num(v, sheet, i, col, errors):
    if v in (None, ""):
        return
    try:
        float(v)
    except (TypeError, ValueError):
        errors.append(f"{sheet} row {i}: '{col}' must be numeric, got {v!r}")

def import_workbook(db: Session, tenant_id: int, filename: str, content: bytes, uploaded_by: str, dry_run=False):
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        return {"ok": False, "errors": ["Not a readable .xlsx file"]}
    data = {n: _rows(wb[n], s, errors) for n, s in SHEETS.items() if n in wb.sheetnames}
    if not data:
        errors.append("Workbook has none of the expected sheets: Positions, KPIs, Records")
    pos_codes = {c for (c,) in db.query(Position.code).filter_by(tenant_id=tenant_id)} | {d["code"] for _, d in data.get("Positions", []) if d.get("code")}
    kpi_codes = {c for (c,) in db.query(KPI.code).filter_by(tenant_id=tenant_id)} | {d["code"] for _, d in data.get("KPIs", []) if d.get("code")}
    for i, d in data.get("KPIs", []):
        if d.get("position_code") and d["position_code"] not in pos_codes:
            errors.append(f"KPIs row {i}: unknown position_code {d['position_code']!r}")
        _num(d.get("target"), "KPIs", i, "target", errors)
        _num(d.get("weight"), "KPIs", i, "weight", errors)
        if d.get("direction") and d["direction"] not in ("higher", "lower"):
            errors.append(f"KPIs row {i}: direction must be 'higher' or 'lower'")
    for i, d in data.get("Records", []):
        if d.get("kpi_code") and d["kpi_code"] not in kpi_codes:
            errors.append(f"Records row {i}: unknown kpi_code {d['kpi_code']!r}")
        _num(d.get("actual"), "Records", i, "actual", errors)
    counts = {k: len(v) for k, v in data.items()}
    if errors or dry_run:
        return {"ok": not errors, "dry_run": dry_run, "errors": errors, "would_import": counts}

    batch = ImportBatch(tenant_id=tenant_id, filename=filename, sha256=hashlib.sha256(content).hexdigest(),
                        uploaded_by=uploaded_by, rows_imported=sum(counts.values()))
    db.add(batch); db.flush()
    for _, d in data.get("Positions", []):
        obj = db.query(Position).filter_by(tenant_id=tenant_id, code=d["code"]).first() or Position(tenant_id=tenant_id, code=d["code"], title="")
        for f in ("title", "tier", "department", "reports_to", "incumbent"):
            if d.get(f) is not None: setattr(obj, f, str(d[f]))
        db.add(obj)
    db.flush()
    for _, d in data.get("KPIs", []):
        obj = db.query(KPI).filter_by(tenant_id=tenant_id, code=d["code"]).first() or KPI(tenant_id=tenant_id, code=d["code"], target=0)
        obj.position_code, obj.kra, obj.name, obj.target = d["position_code"], d["kra"], d["name"], float(d["target"])
        obj.unit = d.get("unit") or "%"; obj.direction = d.get("direction") or "higher"
        obj.weight = float(d["weight"]) if d.get("weight") not in (None, "") else 1.0
        obj.frequency = d.get("frequency") or "monthly"
        db.add(obj)
    db.flush()
    for _, d in data.get("Records", []):
        rec = db.query(Record).filter_by(tenant_id=tenant_id, kpi_code=d["kpi_code"], period=str(d["period"])).first() or Record(tenant_id=tenant_id, kpi_code=d["kpi_code"], period=str(d["period"]))
        rec.actual = float(d["actual"]) if d.get("actual") not in (None, "") else None
        for f in ("evidence_ref", "verifier", "variance_note", "corrective_action", "institutionalized_as"):
            if d.get(f) is not None: setattr(rec, f, str(d[f]))
        rec.status = str(d.get("status") or "OPEN").upper()
        rec.batch_id = batch.id
        db.add(rec)
    db.commit()
    return {"ok": True, "dry_run": False, "errors": [], "imported": counts, "batch_id": batch.id, "sha256": batch.sha256}
