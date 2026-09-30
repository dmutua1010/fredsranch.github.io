"""Creates templates/hmos_import_template.xlsx - the file admins fill in and upload.
Sample rows are generic (no real client data)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import pathlib

wb = Workbook()
sheets = {
    "Positions": (["code", "title", "tier", "department", "reports_to", "incumbent"],
                  [["GM", "General Manager", "Executive", "Management", None, None],
                   ["OPS", "Operations Manager", "Executive", "Operations", "GM", None],
                   ["FNB", "F&B Manager", "Operations", "Food & Beverage", "OPS", None]]),
    "KPIs": (["code", "position_code", "kra", "name", "unit", "direction", "target", "weight", "frequency"],
             [["K-FIN-01", "GM", "Financial Stewardship", "Gross revenue vs budget", "%", "higher", 100, 2, "monthly"],
              ["K-FIN-02", "OPS", "Financial Stewardship", "F&B cost ratio", "%", "lower", 30, 1, "monthly"],
              ["K-SYS-01", "OPS", "Systems & Structures", "SOP compliance audit score", "%", "higher", 95, 1, "quarterly"],
              ["K-GUE-01", "FNB", "Guest Experience & Brand Reputation", "Guest satisfaction index", "%", "higher", 92, 1, "monthly"],
              ["K-CUL-01", "OPS", "SHIELD Culture Adoption", "Culture action plan completion", "%", "higher", 100, 1, "quarterly"]]),
    "Records": (["kpi_code", "period", "actual", "evidence_ref", "verifier", "variance_note", "corrective_action", "institutionalized_as", "status"],
                [["K-FIN-01", "2026-Q3", 97, "ACC-RPT-0930", "Chief Accountant", "Shortfall in banquet bookings", "Corporate sales push", None, "OPEN"],
                 ["K-FIN-02", "2026-Q3", 31, "POS-AUD-Q3", "Chief Accountant", None, None, None, "CLOSED"],
                 ["K-SYS-01", "2026-Q3", 96, None, None, None, None, None, "OPEN"]]),   # no evidence -> INCOMPLETE
}
wb.remove(wb.active)
for name, (hdr, rows) in sheets.items():
    ws = wb.create_sheet(name); ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows: ws.append(r)
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 24
out = pathlib.Path(__file__).parent / "hmos_import_template.xlsx"
wb.save(out); print("wrote", out)
